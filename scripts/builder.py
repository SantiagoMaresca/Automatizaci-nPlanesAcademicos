import yaml
import io
import operator
import re
import docx2txt
import re
from docx import Document


#Lee los campos del Archivo .docx, guarda en un dicionario 'Datos' y genera un .yaml con los datos de interes
#Se utilizan variables Representativas del campo, o en otro caso para tablas las variables: T1row1Col1 (Table 1 row 1 Col 1 )
def generateYaml(file):
    document = Document (file)
    tables = document.tables
    datos={}
    datos[0]=[{'Asignatura': "Análisis visual de datos"}]
    
    #Lee los campos de interes del .docx y genera un diccioario 'datos'
    datos[1]=[{'Carrera':tables[0].cell(1,3).text}]
    datos[2]=[{'PlanEstudios':tables[0].cell(2,3).text}]
   
    datos[3]=[{'NombreEspañol':tables[0].cell(3,3).text}]
    datos[3].append({'NombreIngles': tables[0].cell(4,3).text})
   
    datos[4]=[{'Docente': tables[0].cell(5,3).text}]
    datos[5]=[{'Año':tables[0].cell(6,3).text}]
    datos[6]=[{'Semestre':tables[0].cell(7,3).text}]
    datos[7]= [{'Previas':tables[0].cell(9,3).text}]
    datos[8]=[{'Duracion':tables[0].cell(11,3).text}]
    datos[8].append({'Creditos':tables[0].cell(11,4).text})
    datos[9]=[{'TipoCurso':tables[0].cell(13,3).text}]
    
    #Datos de la tabla Categoria del Curso -Columna : Colocar una Cruz
    datos[10]=[{'T1row1Col1':tables[0].cell(15,3).text }]
    datos[10].append({'T1row2Col1':tables[0].cell(16,3).text})
    datos[10].append({'T1row3Col1':tables[0].cell(17,3).text})
    datos[10].append({'T1row4Col1':tables[0].cell(18,3).text})
    datos[10].append({'T1row5Col1':tables[0].cell(19,3).text})
    datos[10].append({'T1row6Col1':tables[0].cell(20,3).text})

    #Datos de la tabla Categoria del Curso - Columna : Comentarios   
    datos[10].append({'T1row1Col2':tables[0].cell(15,4).text })
    datos[10].append({'T1row2Col2':tables[0].cell(16,4).text})
    datos[10].append({'T1row3Col2':tables[0].cell(17,4).text})
    datos[10].append({'T1row4Col2':tables[0].cell(18,4).text})
    datos[10].append({'T1row5Col2':tables[0].cell(19,4).text})
    datos[10].append({'T1row6Col2':tables[0].cell(20,4).text})

    #Datos de los puntos 11 -14 
    datos[11]=[{'Resumen':tables[0].cell(22,1).text}]
    datos[12]=[{'Objetivos':tables[0].cell(24,1).text}]
    datos[13]=[{'Resultados':tables[0].cell(26,1).text}]
    datos[14]=[{'Contenido':tables[0].cell(28,1).text}]

    #Datos de Métodos didacticos - Columna: Descripción y fila : Presentaciones Magistrales
    datos[15]=[{'T2row1Col2':tables[0].cell(31,3).text}]
    
    #Datos de Métodos didacticos - Columna: Descripción y fila: Resolución de ejercicios
    datos[15].append({'T2row2Col2':tables[0].cell(32,3).text})

    #Datos de Formás de evaluación - Columan Descripción y fila: Resolución de ejercicios
    datos[16]=[{'T3row1Col1':tables[0].cell(35,3).text}]

    #Datos de biblografía básica
    datos[17]=[{'Biblobasica':tables[0].cell(37,1).text}]
    
    #Datos de biblografía ampliatoria
    datos[18]=[{'BibloAmpliatoria':tables[0].cell(39,1).text}]

    #se crea el archivo yaml pasando el diccionario 'datos'
    yamlname=file.replace('docx','yaml')
    with io.open('../files/yamls/'+yamlname, 'w', encoding='ISO 8859-1') as outfile:
        #print(sorted(datos))
        data=sorted(datos.items())
        print(data)
        yaml.dump(data, outfile, default_flow_style=False, allow_unicode=True)
    return datos

    def cargarYaml(archivo):
        with open(archivo, 'r') as stream:
            data_loaded = yaml.safe_load(stream)
        return  data_loaded
 
# ISO 8859-1'
    # with io.open('data2.yaml', 'w', encoding='utf8') as outfile:
    #     yaml.dump(datos, outfile, default_flow_style=False, allow_unicode=True)


#se genera el Yaml para Análisis visual de datos
generateYaml("Analisis_visual_de_datos.docx")
